from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs
from pathlib import Path
import json, os, uuid, re, io, zipfile, shutil, threading, webbrowser, mimetypes
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage

ROOT = Path(__file__).resolve().parent
UPLOADS = ROOT / 'uploads'; OUTPUT = ROOT / 'output'
UPLOADS.mkdir(exist_ok=True); OUTPUT.mkdir(exist_ok=True)
PORT = 8765
STATE = {'pdf': None, 'pages': 0}

HTML = None


def safe_pdf_name(name):
    name = Path(name).name
    if not name.lower().endswith('.pdf'): name += '.pdf'
    return re.sub(r'[^A-Za-z0-9._ -]', '_', name)

def page_rect_from_norm(page, rect):
    x,y,w,h = rect
    return fitz.Rect(x*page.rect.width, y*page.rect.height, (x+w)*page.rect.width, (y+h)*page.rect.height)

def get_words_in_rect(page, rect):
    words = []
    for w in page.get_text('words'):
        r = fitz.Rect(w[:4])
        cx = (r.x0+r.x1)/2.0
        cy = (r.y0+r.y1)/2.0
        if rect_containment_ratio(r, rect) >= 0.90 and rect.contains(fitz.Point(cx, cy)):
            words.append((r, w[4]))
    return words

def build_text_lines(page, rect, text_mode='merge'):
    """
    Build text rows from native PDF words.
    merge  -> original compact line, e.g. 'Product name FANCY'
    split  -> infer a large label/value gap and write
              'Product name : FANCY'
    """
    words = get_words_in_rect(page, rect)
    words.sort(key=lambda z:(z[0].y0, z[0].x0))
    rows=[]
    for r,t in words:
        placed=False
        for row in rows:
            if abs(r.y0-row['y']) <= max(3.0, min(r.height,row['h'])*0.65):
                row['items'].append((r,t))
                row['y']=(row['y']+r.y0)/2.0
                row['h']=max(row['h'],r.height)
                placed=True
                break
        if not placed:
            rows.append({'y':r.y0,'h':r.height,'items':[(r,t)]})

    lines=[]
    for row in rows:
        row['items'].sort(key=lambda z:z[0].x0)
        items=row['items']
        if not items:
            continue

        if text_mode != 'split' or len(items) < 2:
            lines.append(' '.join(t for _,t in items).strip())
            continue

        # Find the largest horizontal gap. In a label/value layout this gap
        # is substantially larger than normal word spacing.
        gaps=[]
        for j in range(len(items)-1):
            a=items[j][0]; b=items[j+1][0]
            gaps.append((max(0.0,b.x0-a.x1),j))

        max_gap, split_at=max(gaps, key=lambda x:x[0])
        normal_gaps=[g for g,_ in gaps if g > 0]
        median_gap=sorted(normal_gaps)[len(normal_gaps)//2] if normal_gaps else 0.0
        # If there is only one gap, it is often the label/value column gap.
        # A 20pt+ gap is enough to distinguish it from ordinary word spacing.
        if len(gaps) == 1:
            threshold = 20.0
        else:
            threshold=max(12.0, median_gap*3.0)

        # A real column separation should be visually significant.
        if max_gap >= threshold:
            left=' '.join(t for _,t in items[:split_at+1]).strip()
            right=' '.join(t for _,t in items[split_at+1:]).strip()
            if left and right:
                lines.append(f'{left} : {right}')
            else:
                lines.append(' '.join(t for _,t in items).strip())
        else:
            lines.append(' '.join(t for _,t in items).strip())

    return [x for x in lines if x]


def intersects(a,b):
    return a.intersects(b) or b.intersects(a)

def text_in_rect(page, rect, text_mode='merge'):
    return build_text_lines(page, rect, text_mode)

def rect_containment_ratio(inner, outer):
    """How much of inner is actually inside outer."""
    inter = inner & outer
    ia = max(0.0, inner.width) * max(0.0, inner.height)
    if ia <= 0:
        return 0.0
    return max(0.0, inter.width) * max(0.0, inter.height) / ia

def image_is_strictly_in_path(image_rect, path_rect):
    """
    A path is a user-defined boundary. An image is selected only when the
    image placement is essentially fully inside that boundary.
    Small PDF coordinate rounding differences are tolerated.
    """
    ratio = rect_containment_ratio(image_rect, path_rect)
    cx = (image_rect.x0 + image_rect.x1) / 2.0
    cy = (image_rect.y0 + image_rect.y1) / 2.0
    center_inside = path_rect.contains(fitz.Point(cx, cy))
    # Require almost all of the image to be inside AND its center to be inside.
    # This prevents a neighboring image from being selected merely because its
    # edge touches/crosses the path.
    if ratio < 0.985 or not center_inside:
        return False

    # Reject suspicious placements that are substantially larger than the path.
    if image_rect.width > path_rect.width * 1.03:
        return False
    if image_rect.height > path_rect.height * 1.03:
        return False
    return True

def extract_images_in_rect(doc,page,rect,outdir,path_no):
    found=[]
    seen_placements=set()
    page_area = page.rect.width * page.rect.height

    for imginfo in page.get_images(full=True):
        xref=imginfo[0]
        try:
            rects=page.get_image_rects(xref)
        except Exception:
            continue

        for ir in rects:
            # IMPORTANT: do not use simple intersection here.
            if not image_is_strictly_in_path(ir, rect):
                continue

            # Ignore page-sized background/decorative placements.
            if page_area and (ir.width * ir.height) / page_area >= 0.90:
                continue

            placement_key = (
                xref,
                round(ir.x0, 3), round(ir.y0, 3),
                round(ir.x1, 3), round(ir.y1, 3)
            )
            if placement_key in seen_placements:
                continue

            try:
                data=doc.extract_image(xref)
                ext=data.get('ext','png') or 'png'
                fn=f'page_{page.number+1:03d}_path_{path_no:03d}_image_{len(found)+1:03d}.{ext}'
                (outdir/fn).write_bytes(data['image'])
                found.append({
                    'file':fn,
                    'bbox':[ir.x0,ir.y0,ir.x1,ir.y1],
                    'xref':xref,
                    'containment_ratio':round(rect_containment_ratio(ir,rect),5)
                })
                seen_placements.add(placement_key)
            except Exception as e:
                found.append({'error':str(e),'xref':xref})

    return found

def word_is_inside_path(word_rect, path_rect):
    # Caption text must itself belong to the user's selected path.
    ratio = rect_containment_ratio(word_rect, path_rect)
    cx = (word_rect.x0 + word_rect.x1) / 2.0
    cy = (word_rect.y0 + word_rect.y1) / 2.0
    return ratio >= 0.90 and path_rect.contains(fitz.Point(cx, cy))

def caption_for_image(page, image_rect, path_rect, texts=None):
    """
    Find a caption immediately above/below an image, but ONLY among text
    objects inside the same user path. This prevents a nearby caption from
    another image/path from being associated accidentally.
    """
    words = page.get_text('words')
    candidates = []

    for w in words:
        r = fitz.Rect(w[:4])
        t = w[4].strip()
        if not t or not word_is_inside_path(r, path_rect):
            continue

        cx = (r.x0 + r.x1) / 2
        # Caption should be horizontally aligned with this image.
        if cx < image_rect.x0 - image_rect.width*0.12 or cx > image_rect.x1 + image_rect.width*0.12:
            continue

        below_limit = max(28, image_rect.height*0.28)
        above_limit = max(28, image_rect.height*0.20)

        if r.y0 >= image_rect.y1 - 2 and r.y0 <= image_rect.y1 + below_limit:
            candidates.append((r.y0-image_rect.y1, r, t))
        elif r.y1 <= image_rect.y0 + 2 and r.y1 >= image_rect.y0 - above_limit:
            candidates.append((image_rect.y0-r.y1+20, r, t))

    if not candidates:
        return None

    candidates.sort(key=lambda x:x[0])
    base = candidates[0][1]
    row_tol = max(3.0, base.height*0.9)
    row = [c for c in candidates if abs(c[1].y0-base.y0) <= row_tol]
    row.sort(key=lambda x:x[1].x0)
    row = [c for c in row if c[1].x1 >= image_rect.x0-3 and c[1].x0 <= image_rect.x1+3]

    caption = ' '.join(c[2] for c in row).strip()
    return caption or None

def make_xlsx(result, outdir):
    wb=Workbook()
    ws=wb.active
    ws.title='Text'
    ws.append(['Path','Page','Text'])
    for c in ws[1]:
        c.font=Font(bold=True)

    for x in result['paths']:
        if x['type'] in ('text','auto','image_text'):
            for line in x.get('text',[]):
                ws.append([x['path'],x['page'],line])

    wi=wb.create_sheet('Images')
    # The Image column contains the actual embedded picture, not just a file path.
    wi.append(['Path','Page','Caption','Gambar','File','Type'])
    for c in wi[1]:
        c.font=Font(bold=True)

    image_rows=[]
    for x in result['paths']:
        for im in x.get('images',[]):
            wi.append([x['path'],x['page'],im.get('caption',''),'',
                       im.get('file',''),x['type']])
            row=wi.max_row
            file_path=outdir/'images'/im.get('file','')
            if file_path.exists() and file_path.is_file():
                try:
                    xl=XLImage(str(file_path))
                    # Keep thumbnails compact but readable.
                    max_w,max_h=130,90
                    if xl.width and xl.height:
                        scale=min(max_w/xl.width,max_h/xl.height,1.0)
                        xl.width=max(1,int(xl.width*scale))
                        xl.height=max(1,int(xl.height*scale))
                    wi.add_image(xl, f'D{row}')
                    wi.row_dimensions[row].height=72
                except Exception:
                    pass
            image_rows.append(row)

    wr=wb.create_sheet('Paths')
    wr.append(['Path','Page','Type','X','Y','W','H'])
    for c in wr[1]:
        c.font=Font(bold=True)
    for x in result['paths']:
        wr.append([x['path'],x['page'],x['type'],*x['rect']])

    for sh in wb.worksheets:
        for col in sh.columns:
            letter=col[0].column_letter
            maxlen=max(len(str(c.value or '')) for c in col)
            sh.column_dimensions[letter].width=min(60,max(12,maxlen+2))
            for c in col:
                c.alignment=Alignment(vertical='top',wrap_text=True)

    # Dedicated image column.
    wi.column_dimensions['D'].width=20
    wi.column_dimensions['C'].width=28
    wi.column_dimensions['E'].width=42
    wi.freeze_panes='A2'
    ws.freeze_panes='A2'
    wr.freeze_panes='A2'

    p=outdir/'extracted.xlsx'
    wb.save(p)
    return p

def extract_project(pdf_path, paths, job, text_mode='merge'):
    jobdir=OUTPUT/job
    if jobdir.exists(): shutil.rmtree(jobdir)
    (jobdir/'text').mkdir(parents=True); (jobdir/'images').mkdir()
    doc=fitz.open(pdf_path); result={'source':Path(pdf_path).name,'text_mode':text_mode,'paths':[]}
    for i,p in enumerate(paths,1):
        page_no=int(p['page']); page=doc[page_no-1]; rect=page_rect_from_norm(page,p['rect']); typ=p['type']
        item={'path':i,'page':page_no,'type':typ,'rect':p['rect'],'text':[],'images':[]}
        if typ in ('text','auto','image_text'):
            item['text']=text_in_rect(page,rect,text_mode)
            (jobdir/'text'/f'page_{page_no:03d}_path_{i:03d}.txt').write_text('\n'.join(item['text']),encoding='utf-8')
        if typ in ('image','auto','image_text'):
            ims=extract_images_in_rect(doc,page,rect,jobdir/'images',i)
            for im in ims:
                if 'file' in im:
                    ir=fitz.Rect(im['bbox']); cap=caption_for_image(page,ir,rect,item['text']) if typ in ('image_text','auto') else None
                    if cap:
                        ext=Path(im['file']).suffix; new=re.sub(r'[^A-Za-z0-9._ -]+','_',cap).strip()[:120]+ext
                        src=jobdir/'images'/im['file']; dst=jobdir/'images'/new
                        if dst.exists():
                            k=2
                            while (jobdir/'images'/f'{dst.stem}_{k}{ext}').exists(): k+=1
                            dst=jobdir/'images'/f'{dst.stem}_{k}{ext}'
                        src.rename(dst); im['file']=dst.name; im['caption']=cap
                    else: im['caption']=''
            item['images']=ims
        result['paths'].append(item)
    doc.close()
    (jobdir/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    xlsx=make_xlsx(result,jobdir)
    zipfile_path=OUTPUT/f'{job}.zip'
    if zipfile_path.exists(): zipfile_path.unlink()
    with zipfile.ZipFile(zipfile_path,'w',zipfile.ZIP_DEFLATED) as z:
        for f in jobdir.rglob('*'):
            if f.is_file(): z.write(f,f.relative_to(jobdir))
    return result, zipfile_path

class Handler(BaseHTTPRequestHandler):
    def send_json(self,obj,status=200):
        data=json.dumps(obj,ensure_ascii=False).encode(); self.send_response(status); self.send_header('Content-Type','application/json'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        u=urlparse(self.path); q=parse_qs(u.query)
        if u.path=='/':
            p = ROOT / 'templates' / 'index.html'
            data = p.read_bytes()
            self.send_response(200)
            self.send_header('Content-Type','text/html; charset=utf-8')
            self.send_header('Cache-Control','no-cache')
            self.send_header('Content-Length',str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return

        if u.path.startswith('/static/'):
            rel = u.path[len('/static/'):].lstrip('/')
            p = (ROOT / 'static' / rel).resolve()
            static_root = (ROOT / 'static').resolve()
            if not str(p).startswith(str(static_root)) or not p.is_file():
                self.send_response(404); self.end_headers(); return
            data = p.read_bytes()
            ctype = mimetypes.guess_type(str(p))[0] or 'application/octet-stream'
            self.send_response(200)
            self.send_header('Content-Type',ctype)
            self.send_header('Cache-Control','no-cache')
            self.send_header('Content-Length',str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return

        if u.path=='/api/page':
            pdf=q.get('pdf',[STATE['pdf']])[0]; page_no=int(q.get('page',['1'])[0]); p=UPLOADS/pdf
            try:
                doc=fitz.open(p); page=doc[page_no-1]; pix=page.get_pixmap(matrix=fitz.Matrix(1.6,1.6),alpha=False); data=pix.tobytes('png'); doc.close()
                self.send_response(200); self.send_header('Content-Type','image/png'); self.send_header('Cache-Control','no-cache'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data)
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        if u.path.startswith('/download/'):
            name=Path(u.path.split('/download/',1)[1]).name; p=OUTPUT/name
            if not p.exists(): self.send_json({'error':'not found'},404); return
            data=p.read_bytes(); self.send_response(200); self.send_header('Content-Type','application/octet-stream'); self.send_header('Content-Disposition',f'attachment; filename="{name}"'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data); return
        self.send_response(404); self.end_headers()
    def do_POST(self):
        if self.path=='/api/upload':
            n=int(self.headers.get('Content-Length','0')); data=self.rfile.read(n); name=safe_pdf_name(__import__('urllib.parse').parse.unquote(self.headers.get('X-Filename','document.pdf'))); path=UPLOADS/name; path.write_bytes(data)
            try:
                doc=fitz.open(path); pages=len(doc); doc.close(); STATE.update(pdf=name,pages=pages); self.send_json({'pdf':name,'pages':pages})
            except Exception as e: path.unlink(missing_ok=True); self.send_json({'error':str(e)},400)
            return
        if self.path=='/api/extract':
            try:
                n=int(self.headers.get('Content-Length','0')); body=json.loads(self.rfile.read(n)); pdf=body['pdf']; paths=body['paths']; text_mode=body.get('text_mode','merge');
                if not paths: raise ValueError('No paths defined')
                job='extract_'+uuid.uuid4().hex[:10]; result,z=extract_project(UPLOADS/pdf,paths,job,text_mode); summary={'source':result['source'],'paths':len(paths),'text_lines':sum(len(x['text']) for x in result['paths']),'images':sum(len(x['images']) for x in result['paths']),'job':job,'download':'/download/'+z.name}
                self.send_json({'summary':summary,'download':'/download/'+z.name})
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        self.send_response(404); self.end_headers()

if __name__=='__main__':
    server=ThreadingHTTPServer(('127.0.0.1',PORT),Handler)
    url=f'http://127.0.0.1:{PORT}/'
    print('PDF Path Extractor:',url)
    threading.Timer(0.8,lambda:webbrowser.open(url)).start()
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close()
