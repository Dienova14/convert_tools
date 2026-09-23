from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs, unquote
from pathlib import Path
import json, os, uuid, re, io, zipfile, shutil, threading, webbrowser, mimetypes, hashlib
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage

ROOT = Path(__file__).resolve().parent
UPLOADS = ROOT / 'uploads'; OUTPUT = ROOT / 'output'
UPLOADS.mkdir(exist_ok=True); OUTPUT.mkdir(exist_ok=True)
PORT = 8765
# STATE tracks the "last active" pdf purely as a convenience default for
# endpoints that don't receive an explicit pdf name. Every endpoint that
# matters (page render, extract, batch extract, delete) takes an explicit
# pdf name from the client, so multiple PDFs can be worked with side by side.
STATE = {'pdf': None, 'pages': 0}

HTML = None


def safe_pdf_name(name):
    name = Path(name).name
    if not name.lower().endswith('.pdf'): name += '.pdf'
    return re.sub(r'[^A-Za-z0-9._ -]', '_', name)


def unique_pdf_path(name, data):
    """Return the path to store this PDF under.

    Re-uploading bytes identical to an existing file reuses that file's name
    (so re-opening the same PDF doesn't create endless duplicates). A
    different file that happens to share a name gets a numeric suffix instead
    of silently overwriting the existing one, which used to corrupt any
    Paths already defined against the original file.
    """
    path = UPLOADS / name
    new_hash = hashlib.sha1(data).hexdigest()
    if not path.exists():
        return path
    if hashlib.sha1(path.read_bytes()).hexdigest() == new_hash:
        return path
    stem, suffix = path.stem, path.suffix
    i = 2
    while True:
        candidate = UPLOADS / f'{stem}_{i}{suffix}'
        if not candidate.exists():
            return candidate
        if hashlib.sha1(candidate.read_bytes()).hexdigest() == new_hash:
            return candidate
        i += 1


def list_pdfs():
    result = []
    for p in UPLOADS.glob('*.pdf'):
        try:
            doc = fitz.open(p); pages = len(doc); doc.close()
        except Exception:
            pages = 0
        st = p.stat()
        result.append({'name': p.name, 'pages': pages, 'size': st.st_size, 'mtime': st.st_mtime})
    result.sort(key=lambda x: x['mtime'], reverse=True)
    return result


def page_rect_from_norm(page, rect):
    x,y,w,h = rect
    return fitz.Rect(x*page.rect.width, y*page.rect.height, (x+w)*page.rect.width, (y+h)*page.rect.height)

def get_words_in_rect(page, rect):
    words = []
    for w in page.get_text('words'):
        r = fitz.Rect(w[:4])
        cx = (r.x0+r.x1)/2.0
        cy = (r.y0+r.y1)/2.0
        if rect_containment_ratio(r, rect) >= 0.90 and rect.contains(fitz.Point(cx, cy)):
            words.append((r, w[4]))
    return words

def build_text_lines(page, rect, text_mode='merge'):
    """
    Build text rows from native PDF words.
    merge  -> original compact line, e.g. 'Product name FANCY'
    split  -> infer a large label/value gap and write
              'Product name : FANCY'
    """
    words = get_words_in_rect(page, rect)
    words.sort(key=lambda z:(z[0].y0, z[0].x0))
    rows=[]
    for r,t in words:
        placed=False
        for row in rows:
            if abs(r.y0-row['y']) <= max(3.0, min(r.height,row['h'])*0.65):
                row['items'].append((r,t))
                row['y']=(row['y']+r.y0)/2.0
                row['h']=max(row['h'],r.height)
                placed=True
                break
        if not placed:
            rows.append({'y':r.y0,'h':r.height,'items':[(r,t)]})

    lines=[]
    for row in rows:
        row['items'].sort(key=lambda z:z[0].x0)
        items=row['items']
        if not items:
            continue

        if text_mode != 'split' or len(items) < 2:
            lines.append(' '.join(t for _,t in items).strip())
            continue

        # Find the largest horizontal gap. In a label/value layout this gap
        # is substantially larger than normal word spacing.
        gaps=[]
        for j in range(len(items)-1):
            a=items[j][0]; b=items[j+1][0]
            gaps.append((max(0.0,b.x0-a.x1),j))

        max_gap, split_at=max(gaps, key=lambda x:x[0])
        normal_gaps=[g for g,_ in gaps if g > 0]
        median_gap=sorted(normal_gaps)[len(normal_gaps)//2] if normal_gaps else 0.0
        # If there is only one gap, it is often the label/value column gap.
        # A 20pt+ gap is enough to distinguish it from ordinary word spacing.
        if len(gaps) == 1:
            threshold = 20.0
        else:
            threshold=max(12.0, median_gap*3.0)

        # A real column separation should be visually significant.
        if max_gap >= threshold:
            left=' '.join(t for _,t in items[:split_at+1]).strip()
            right=' '.join(t for _,t in items[split_at+1:]).strip()
            if left and right:
                lines.append(f'{left} : {right}')
            else:
                lines.append(' '.join(t for _,t in items).strip())
        else:
            lines.append(' '.join(t for _,t in items).strip())

    return [x for x in lines if x]


def intersects(a,b):
    return a.intersects(b) or b.intersects(a)

def text_in_rect(page, rect, text_mode='merge'):
    return build_text_lines(page, rect, text_mode)


def ocr_fallback_text(page, rect, dpi=300):
    """Best-effort OCR for a path region when native text extraction finds
    nothing (typically a scanned/image-only PDF page). Silently returns []
    when pytesseract/Pillow aren't installed, or when OCR itself fails —
    this must never break extraction for people without OCR set up.
    """
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return []
    try:
        zoom = dpi / 72.0
        pix = page.get_pixmap(clip=rect, matrix=fitz.Matrix(zoom, zoom), alpha=False)
        img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
        text = pytesseract.image_to_string(img, lang='ind+eng')
    except Exception:
        try:
            text = pytesseract.image_to_string(img)
        except Exception:
            return []
    return [l.strip() for l in text.splitlines() if l.strip()]

def rect_containment_ratio(inner, outer):
    """How much of inner is actually inside outer."""
    inter = inner & outer
    ia = max(0.0, inner.width) * max(0.0, inner.height)
    if ia <= 0:
        return 0.0
    return max(0.0, inter.width) * max(0.0, inter.height) / ia

def image_is_strictly_in_path(image_rect, path_rect):
    """
    A path is a user-defined boundary. An image is selected only when the
    image placement is essentially fully inside that boundary.
    Small PDF coordinate rounding differences are tolerated.
    """
    ratio = rect_containment_ratio(image_rect, path_rect)
    cx = (image_rect.x0 + image_rect.x1) / 2.0
    cy = (image_rect.y0 + image_rect.y1) / 2.0
    center_inside = path_rect.contains(fitz.Point(cx, cy))
    # Require almost all of the image to be inside AND its center to be inside.
    # This prevents a neighboring image from being selected merely because its
    # edge touches/crosses the path.
    if ratio < 0.985 or not center_inside:
        return False

    # Reject suspicious placements that are substantially larger than the path.
    if image_rect.width > path_rect.width * 1.03:
        return False
    if image_rect.height > path_rect.height * 1.03:
        return False
    return True

def extract_images_in_rect(doc,page,rect,outdir,path_no):
    found=[]
    seen_placements=set()
    page_area = page.rect.width * page.rect.height

    for imginfo in page.get_images(full=True):
        xref=imginfo[0]
        try:
            rects=page.get_image_rects(xref)
        except Exception:
            continue

        for ir in rects:
            # IMPORTANT: do not use simple intersection here.
            if not image_is_strictly_in_path(ir, rect):
                continue

            # Ignore page-sized background/decorative placements.
            if page_area and (ir.width * ir.height) / page_area >= 0.90:
                continue

            placement_key = (
                xref,
                round(ir.x0, 3), round(ir.y0, 3),
                round(ir.x1, 3), round(ir.y1, 3)
            )
            if placement_key in seen_placements:
                continue

            try:
                data=doc.extract_image(xref)
                ext=data.get('ext','png') or 'png'
                fn=f'page_{page.number+1:03d}_path_{path_no:03d}_image_{len(found)+1:03d}.{ext}'
                (outdir/fn).write_bytes(data['image'])
                found.append({
                    'file':fn,
                    'bbox':[ir.x0,ir.y0,ir.x1,ir.y1],
                    'xref':xref,
                    'containment_ratio':round(rect_containment_ratio(ir,rect),5)
                })
                seen_placements.add(placement_key)
            except Exception as e:
                found.append({'error':str(e),'xref':xref})

    return found

def word_is_inside_path(word_rect, path_rect):
    # Caption text must itself belong to the user's selected path.
    ratio = rect_containment_ratio(word_rect, path_rect)
    cx = (word_rect.x0 + word_rect.x1) / 2.0
    cy = (word_rect.y0 + word_rect.y1) / 2.0
    return ratio >= 0.90 and path_rect.contains(fitz.Point(cx, cy))

def caption_for_image(page, image_rect, path_rect, texts=None):
    """
    Find a caption immediately above/below an image, but ONLY among text
    objects inside the same user path. This prevents a nearby caption from
    another image/path from being associated accidentally.
    """
    words = page.get_text('words')
    candidates = []

    for w in words:
        r = fitz.Rect(w[:4])
        t = w[4].strip()
        if not t or not word_is_inside_path(r, path_rect):
            continue

        cx = (r.x0 + r.x1) / 2
        # Caption should be horizontally aligned with this image.
        if cx < image_rect.x0 - image_rect.width*0.12 or cx > image_rect.x1 + image_rect.width*0.12:
            continue

        below_limit = max(28, image_rect.height*0.28)
        above_limit = max(28, image_rect.height*0.20)

        if r.y0 >= image_rect.y1 - 2 and r.y0 <= image_rect.y1 + below_limit:
            candidates.append((r.y0-image_rect.y1, r, t))
        elif r.y1 <= image_rect.y0 + 2 and r.y1 >= image_rect.y0 - above_limit:
            candidates.append((image_rect.y0-r.y1+20, r, t))

    if not candidates:
        return None

    candidates.sort(key=lambda x:x[0])
    base = candidates[0][1]
    row_tol = max(3.0, base.height*0.9)
    row = [c for c in candidates if abs(c[1].y0-base.y0) <= row_tol]
    row.sort(key=lambda x:x[1].x0)
    row = [c for c in row if c[1].x1 >= image_rect.x0-3 and c[1].x0 <= image_rect.x1+3]

    caption = ' '.join(c[2] for c in row).strip()
    return caption or None

def _autosize_and_freeze(wb):
    for sh in wb.worksheets:
        for col in sh.columns:
            letter=col[0].column_letter
            maxlen=max((len(str(c.value or '')) for c in col), default=10)
            sh.column_dimensions[letter].width=min(60,max(12,maxlen+2))
            for c in col:
                c.alignment=Alignment(vertical='top',wrap_text=True)
        sh.freeze_panes='A2'

def make_xlsx(result, outdir):
    wb=Workbook()
    ws=wb.active
    ws.title='Text'
    ws.append(['Path','Page','Text'])
    for c in ws[1]:
        c.font=Font(bold=True)

    for x in result['paths']:
        if x['type'] in ('text','auto','image_text'):
            for line in x.get('text',[]):
                ws.append([x['path'],x['page'],line])

    wi=wb.create_sheet('Images')
    # The Image column contains the actual embedded picture, not just a file path.
    wi.append(['Path','Page','Caption','Gambar','File','Type'])
    for c in wi[1]:
        c.font=Font(bold=True)

    for x in result['paths']:
        for im in x.get('images',[]):
            wi.append([x['path'],x['page'],im.get('caption',''),'',
                       im.get('file',''),x['type']])
            row=wi.max_row
            file_path=outdir/'images'/im.get('file','')
            if file_path.exists() and file_path.is_file():
                try:
                    xl=XLImage(str(file_path))
                    # Keep thumbnails compact but readable.
                    max_w,max_h=130,90
                    if xl.width and xl.height:
                        scale=min(max_w/xl.width,max_h/xl.height,1.0)
                        xl.width=max(1,int(xl.width*scale))
                        xl.height=max(1,int(xl.height*scale))
                    wi.add_image(xl, f'D{row}')
                    wi.row_dimensions[row].height=72
                except Exception:
                    pass

    wr=wb.create_sheet('Paths')
    wr.append(['Path','Page','Type','X','Y','W','H'])
    for c in wr[1]:
        c.font=Font(bold=True)
    for x in result['paths']:
        wr.append([x['path'],x['page'],x['type'],*x['rect']])

    _autosize_and_freeze(wb)
    wi.column_dimensions['D'].width=20
    wi.column_dimensions['C'].width=28
    wi.column_dimensions['E'].width=42

    p=outdir/'extracted.xlsx'
    wb.save(p)
    return p


def make_combined_xlsx(combined, outdir):
    """Same shape as make_xlsx but for a batch of PDFs: adds a 'Source PDF'
    column and a dedicated 'Errors' sheet for files that failed to process,
    so one bad/mismatched file in the batch doesn't hide the rest."""
    wb=Workbook()
    ws=wb.active; ws.title='Text'
    ws.append(['Source PDF','Path','Page','Text'])
    for c in ws[1]: c.font=Font(bold=True)

    wi=wb.create_sheet('Images')
    wi.append(['Source PDF','Path','Page','Caption','Gambar','File','Type'])
    for c in wi[1]: c.font=Font(bold=True)

    we=wb.create_sheet('Errors')
    we.append(['Source PDF','Error'])
    for c in we[1]: c.font=Font(bold=True)

    for f in combined['files']:
        if 'error' in f:
            we.append([f.get('source',''), f['error']])
            continue
        source=f['source']
        for x in f['paths']:
            if x['type'] in ('text','auto','image_text'):
                for line in x.get('text',[]):
                    ws.append([source,x['path'],x['page'],line])
            for im in x.get('images',[]):
                wi.append([source,x['path'],x['page'],im.get('caption',''),'',
                           im.get('file',''),x['type']])
                row=wi.max_row
                file_path=outdir/Path(source).stem/'images'/im.get('file','')
                if file_path.exists() and file_path.is_file():
                    try:
                        xl=XLImage(str(file_path))
                        max_w,max_h=130,90
                        if xl.width and xl.height:
                            scale=min(max_w/xl.width,max_h/xl.height,1.0)
                            xl.width=max(1,int(xl.width*scale))
                            xl.height=max(1,int(xl.height*scale))
                        wi.add_image(xl, f'E{row}')
                        wi.row_dimensions[row].height=72
                    except Exception:
                        pass

    _autosize_and_freeze(wb)
    wi.column_dimensions['E'].width=20
    wi.column_dimensions['D'].width=28
    wi.column_dimensions['F'].width=42

    p=outdir/'combined_extracted.xlsx'
    wb.save(p)
    return p



def _rect_union(rects):
    """Return the smallest rectangle containing all supplied rectangles."""
    rects=[fitz.Rect(r) for r in rects if r and r.width>0 and r.height>0]
    if not rects:
        return None
    out=fitz.Rect(rects[0])
    for r in rects[1:]:
        out |= r
    return out

def _word_rect(w):
    return fitz.Rect(float(w[0]),float(w[1]),float(w[2]),float(w[3]))

def _overlap_ratio(a,b):
    inter=a & b
    if inter.is_empty:
        return 0.0
    return (inter.width*inter.height)/max(1e-9,a.width*a.height)

def _horizontal_overlap(a,b):
    overlap=max(0.0,min(a.x1,b.x1)-max(a.x0,b.x0))
    return overlap/max(1e-9,min(a.width,b.width))

def _associate_word_to_image(word_rect, image_rect, page_rect):
    """Score whether a text word belongs to an image/caption region.

    A word is considered attached when it is inside/overlapping the image, or
    is a nearby caption with meaningful horizontal overlap.  The score is
    deliberately conservative so one wide caption line is split between the
    actual images instead of being attached to every image on the page.
    """
    if _overlap_ratio(word_rect,image_rect) >= 0.20:
        return 100.0 + _overlap_ratio(word_rect,image_rect)*20.0
    if image_rect.contains(fitz.Point((word_rect.x0+word_rect.x1)/2,(word_rect.y0+word_rect.y1)/2)):
        return 95.0

    gap=0.0
    if word_rect.y1 < image_rect.y0:
        gap=image_rect.y0-word_rect.y1
    elif image_rect.y1 < word_rect.y0:
        gap=word_rect.y0-image_rect.y1
    else:
        gap=0.0
    if gap > max(30.0,page_rect.height*0.035):
        return -1.0

    hov=_horizontal_overlap(word_rect,image_rect)
    if hov < 0.18:
        return -1.0
    # Prefer captions directly adjacent to the image and with stronger x overlap.
    return 50.0 + hov*25.0 - gap*0.8

def _pad_rect(rect,page):
    px=min(5.0,page.rect.width*0.008)
    py=min(5.0,page.rect.height*0.008)
    return fitz.Rect(
        max(page.rect.x0,rect.x0-px),
        max(page.rect.y0,rect.y0-py),
        min(page.rect.x1,rect.x1+px),
        min(page.rect.y1,rect.y1+py),
    )

def generate_paths_for_pdf(pdf_path):
    """Generate editable paths from actual PDF objects.

    The previous generator merged nearby objects too aggressively, which could
    turn an entire page into one `auto` path.  This version creates a path per
    embedded image and a path per remaining native-text block.  Text located
    inside/next to an image is attached to that image, so the generated type is
    automatically `image_text`, `image`, or `text` — never `auto`.
    """
    doc=fitz.open(pdf_path)
    generated=[]

    for page_no,page in enumerate(doc,1):
        page_rect=page.rect
        page_area=max(1.0,page_rect.width*page_rect.height)

        # Native text words retain their original block number.  This lets us
        # keep normal paragraphs/tables together while still splitting a wide
        # caption line across separate images when necessary.
        raw_words=page.get_text('words') or []
        words=[]
        for w in raw_words:
            if len(w)<5: continue
            text=str(w[4] or '').strip()
            if not text: continue
            words.append({
                'rect':_word_rect(w),
                'text':text,
                'block':int(w[5]) if len(w)>5 else -1,
                'line':int(w[6]) if len(w)>6 else -1,
                'word':int(w[7]) if len(w)>7 else -1,
                'used':False,
            })

        # Collect each visible embedded-image placement, not only unique xrefs.
        images=[]
        seen=[]
        for info in page.get_images(full=True):
            xref=info[0]
            try:
                rects=page.get_image_rects(xref)
            except Exception:
                rects=[]
            for rr in rects:
                r=fitz.Rect(rr)
                if r.width<3 or r.height<3: continue
                # Ignore full-page/background images; they are normally page
                # backgrounds rather than a user-extractable content image.
                if (r.width*r.height)/page_area >= 0.88:
                    continue
                duplicate=False
                for old in seen:
                    if old['xref']==xref and _overlap_ratio(r,old['rect'])>0.98:
                        duplicate=True; break
                if duplicate: continue
                item={'xref':xref,'rect':r,'words':[]}
                images.append(item); seen.append(item)

        # Assign each word to at most one image.  This is especially important
        # for a caption block that spans several product photos.
        for wi,w in enumerate(words):
            candidates=[]
            for ii,im in enumerate(images):
                score=_associate_word_to_image(w['rect'],im['rect'],page_rect)
                if score>=0:
                    candidates.append((score,ii))
            if candidates:
                candidates.sort(reverse=True)
                score,ii=candidates[0]
                # A distant/weak candidate is not enough to steal text from a
                # normal text block.
                if score>=48.0:
                    images[ii]['words'].append(wi)
                    w['used']=True

        # Image paths first.  The union includes attached text/captions.
        for im in images:
            rects=[im['rect']]+[words[i]['rect'] for i in im['words']]
            texts=[words[i]['text'] for i in sorted(im['words'],key=lambda i:(words[i]['rect'].y0,words[i]['rect'].x0))]
            # Tiny standalone image objects are commonly arrows, bullets,
            # masks or decorative PDF fragments. Keep them only when text is
            # actually associated with them; this avoids dozens of useless
            # paths while preserving small labelled icons.
            image_area_ratio=(im['rect'].width*im['rect'].height)/page_area
            if not texts and image_area_ratio < 0.0025:
                continue
            union=_pad_rect(_rect_union(rects),page)
            typ='image_text' if texts else 'image'
            generated.append({
                'page':page_no,
                'type':typ,
                'rect':[union.x0/page_rect.width,union.y0/page_rect.height,union.width/page_rect.width,union.height/page_rect.height],
            })

        # Remaining native text stays grouped by its original PDF text block.
        # This prevents individual words/lines from becoming hundreds of paths.
        blocks={}
        for i,w in enumerate(words):
            if w['used']: continue
            blocks.setdefault(w['block'],[]).append(i)

        text_regions=[]
        for block_no,idxs in blocks.items():
            if not idxs: continue
            rect=_rect_union([words[i]['rect'] for i in idxs])
            if not rect or rect.width<2 or rect.height<2: continue
            text_regions.append((rect,idxs))

        # Some PDFs report one large block for several disconnected text
        # is clearly larger than normal line spacing.
        split_regions=[]
        for rect,idxs in text_regions:
            idxs=sorted(idxs,key=lambda i:(words[i]['rect'].y0,words[i]['rect'].x0))
            groups=[]; current=[]; last_y=None; last_h=None
            for i in idxs:
                wr=words[i]['rect']
                if not current:
                    current=[i]; last_y=wr.y0; last_h=wr.height; continue
                gap=wr.y0-(last_y+last_h)
                if gap<=max(8.0,last_h*0.9):
                    current.append(i)
                else:
                    groups.append(current); current=[i]
                last_y=wr.y0; last_h=wr.height
            if current: groups.append(current)
            # Only split when there is a genuinely large vertical separation;
            # otherwise retain the original PDF text block.
            if len(groups)>1 and rect.height > 2*max(g[0] and words[g[0]]['rect'].height or 1 for g in groups):
                for g in groups:
                    split_regions.append((_rect_union([words[i]['rect'] for i in g]),g))
            else:
                split_regions.append((rect,idxs))

        for rect,idxs in split_regions:
            union=_pad_rect(rect,page)
            generated.append({
                'page':page_no,
                'type':'text',
                'rect':[union.x0/page_rect.width,union.y0/page_rect.height,union.width/page_rect.width,union.height/page_rect.height],
            })

    doc.close()
    # Reading order: top-to-bottom, then left-to-right.
    generated.sort(key=lambda p:(p['page'],round(p['rect'][1]*1000/5)*5,p['rect'][0]))
    return generated

def extract_project(pdf_path, paths, job, text_mode='merge'):
    jobdir=OUTPUT/job
    if jobdir.exists(): shutil.rmtree(jobdir)
    (jobdir/'text').mkdir(parents=True); (jobdir/'images').mkdir()
    doc=fitz.open(pdf_path); result={'source':Path(pdf_path).name,'text_mode':text_mode,'paths':[]}
    for i,p in enumerate(paths,1):
        page_no=int(p['page'])
        if page_no < 1 or page_no > len(doc):
            # Skip paths that don't exist on this PDF (common in batch mode
            # when a template was drawn against a longer/shorter document).
            result['paths'].append({'path':i,'page':page_no,'type':p['type'],'rect':p['rect'],
                                     'text':[],'images':[],'error':'Halaman tidak ada di PDF ini'})
            continue
        page=doc[page_no-1]; rect=page_rect_from_norm(page,p['rect']); typ=p['type']
        item={'path':i,'page':page_no,'type':typ,'rect':p['rect'],'text':[],'images':[]}
        if typ in ('text','auto','image_text'):
            item['text']=text_in_rect(page,rect,text_mode)
            if not item['text']:
                ocr_lines=ocr_fallback_text(page,rect)
                if ocr_lines:
                    item['text']=ocr_lines
                    item['ocr']=True
            (jobdir/'text'/f'page_{page_no:03d}_path_{i:03d}.txt').write_text('\n'.join(item['text']),encoding='utf-8')
        if typ in ('image','auto','image_text'):
            ims=extract_images_in_rect(doc,page,rect,jobdir/'images',i)
            for im in ims:
                if 'file' in im:
                    ir=fitz.Rect(im['bbox']); cap=caption_for_image(page,ir,rect,item['text']) if typ in ('image_text','auto') else None
                    if cap:
                        ext=Path(im['file']).suffix; new=re.sub(r'[^A-Za-z0-9._ -]+','_',cap).strip()[:120]+ext
                        src=jobdir/'images'/im['file']; dst=jobdir/'images'/new
                        if dst.exists():
                            k=2
                            while (jobdir/'images'/f'{dst.stem}_{k}{ext}').exists(): k+=1
                            dst=jobdir/'images'/f'{dst.stem}_{k}{ext}'
                        src.rename(dst); im['file']=dst.name; im['caption']=cap
                    else: im['caption']=''
            item['images']=ims
        result['paths'].append(item)
    doc.close()
    (jobdir/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    xlsx=make_xlsx(result,jobdir)
    zipfile_path=OUTPUT/f'{job}.zip'
    if zipfile_path.exists(): zipfile_path.unlink()
    with zipfile.ZipFile(zipfile_path,'w',zipfile.ZIP_DEFLATED) as z:
        for f in jobdir.rglob('*'):
            if f.is_file(): z.write(f,f.relative_to(jobdir))
    return result, zipfile_path


def batch_extract(pdf_names, paths, text_mode, batch_job):
    """Apply the same set of Paths (drawn once, normalized 0..1 per page) to
    many uploaded PDFs. Intended for PDFs that share a layout (e.g. the same
    invoice/form template) — a path's rect means the same physical spot on
    every file it's applied to. One failing/mismatched file is recorded in
    the batch's Errors sheet instead of aborting the whole run.
    """
    batchdir=OUTPUT/batch_job
    if batchdir.exists(): shutil.rmtree(batchdir)
    batchdir.mkdir(parents=True)
    combined={'text_mode':text_mode,'files':[]}

    for name in pdf_names:
        pdf_path=UPLOADS/name
        if not pdf_path.exists():
            combined['files'].append({'source':name,'error':'File tidak ditemukan di server'})
            continue
        sub_job=f'{batch_job}__{uuid.uuid4().hex[:6]}_{Path(name).stem}'
        try:
            result,_=extract_project(pdf_path,paths,sub_job,text_mode)
        except Exception as e:
            combined['files'].append({'source':name,'error':str(e)})
            continue
        src_dir=OUTPUT/sub_job
        dst_dir=batchdir/Path(name).stem
        if dst_dir.exists(): shutil.rmtree(dst_dir)
        shutil.move(str(src_dir),str(dst_dir))
        per_zip=OUTPUT/f'{sub_job}.zip'
        if per_zip.exists(): per_zip.unlink()
        combined['files'].append(result)

    make_combined_xlsx(combined,batchdir)
    (batchdir/'batch_result.json').write_text(json.dumps(combined,ensure_ascii=False,indent=2),encoding='utf-8')
    zip_path=OUTPUT/f'{batch_job}.zip'
    if zip_path.exists(): zip_path.unlink()
    with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
        for f in batchdir.rglob('*'):
            if f.is_file(): z.write(f,f.relative_to(batchdir))
    return combined, zip_path

class Handler(BaseHTTPRequestHandler):
    def send_json(self,obj,status=200):
        data=json.dumps(obj,ensure_ascii=False).encode(); self.send_response(status); self.send_header('Content-Type','application/json'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        u=urlparse(self.path); q=parse_qs(u.query)
        if u.path=='/':
            p = ROOT / 'templates' / 'index.html'
            data = p.read_bytes()
            self.send_response(200)
            self.send_header('Content-Type','text/html; charset=utf-8')
            self.send_header('Cache-Control','no-cache')
            self.send_header('Content-Length',str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return

        if u.path.startswith('/static/'):
            rel = u.path[len('/static/'):].lstrip('/')
            p = (ROOT / 'static' / rel).resolve()
            static_root = (ROOT / 'static').resolve()
            if not str(p).startswith(str(static_root)) or not p.is_file():
                self.send_response(404); self.end_headers(); return
            data = p.read_bytes()
            ctype = mimetypes.guess_type(str(p))[0] or 'application/octet-stream'
            self.send_response(200)
            self.send_header('Content-Type',ctype)
            self.send_header('Cache-Control','no-cache')
            self.send_header('Content-Length',str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return

        if u.path=='/api/pdfs':
            self.send_json({'pdfs': list_pdfs()})
            return

        if u.path=='/api/page':
            pdf=q.get('pdf',[STATE['pdf']])[0]; page_no=int(q.get('page',['1'])[0]); p=UPLOADS/pdf
            try:
                doc=fitz.open(p); page=doc[page_no-1]; pix=page.get_pixmap(matrix=fitz.Matrix(1.6,1.6),alpha=False); data=pix.tobytes('png'); doc.close()
                self.send_response(200); self.send_header('Content-Type','image/png'); self.send_header('Cache-Control','no-cache'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data)
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        if u.path.startswith('/download/'):
            name=Path(u.path.split('/download/',1)[1]).name; p=OUTPUT/name
            if not p.exists(): self.send_json({'error':'not found'},404); return
            data=p.read_bytes(); self.send_response(200); self.send_header('Content-Type','application/octet-stream'); self.send_header('Content-Disposition',f'attachment; filename="{name}"'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data); return
        self.send_response(404); self.end_headers()
    def do_POST(self):
        if self.path=='/api/upload':
            n=int(self.headers.get('Content-Length','0')); data=self.rfile.read(n)
            name=safe_pdf_name(unquote(self.headers.get('X-Filename','document.pdf')))
            path=unique_pdf_path(name,data)
            is_new=not path.exists()
            if is_new:
                path.write_bytes(data)
            try:
                doc=fitz.open(path)
                if doc.is_encrypted and not doc.authenticate(''):
                    doc.close()
                    if is_new: path.unlink(missing_ok=True)
                    self.send_json({'error':'PDF terkunci dengan password dan tidak bisa dibuka otomatis.'},400)
                    return
                pages=len(doc); doc.close()
                STATE.update(pdf=path.name,pages=pages)
                self.send_json({'pdf':path.name,'pages':pages})
            except Exception as e:
                if is_new: path.unlink(missing_ok=True)
                self.send_json({'error':str(e)},400)
            return
        if self.path=='/api/delete_pdf':
            try:
                n=int(self.headers.get('Content-Length','0')); body=json.loads(self.rfile.read(n))
                name=Path(body.get('pdf','')).name
                p=UPLOADS/name
                if p.exists(): p.unlink()
                if STATE.get('pdf')==name: STATE.update(pdf=None,pages=0)
                self.send_json({'deleted':name})
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        if self.path=='/api/generate_paths':
            try:
                n=int(self.headers.get('Content-Length','0')); body=json.loads(self.rfile.read(n)); pdf=body['pdf']
                if not pdf or not (UPLOADS/pdf).exists(): raise ValueError('PDF tidak ditemukan')
                generated=generate_paths_for_pdf(UPLOADS/pdf)
                self.send_json({'pdf':pdf,'paths':generated,'count':len(generated)})
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        if self.path=='/api/extract':
            try:
                n=int(self.headers.get('Content-Length','0')); body=json.loads(self.rfile.read(n)); pdf=body['pdf']; paths=body['paths']; text_mode=body.get('text_mode','merge');
                if not paths: raise ValueError('No paths defined')
                job='extract_'+uuid.uuid4().hex[:10]; result,z=extract_project(UPLOADS/pdf,paths,job,text_mode); summary={'source':result['source'],'paths':len(paths),'text_lines':sum(len(x['text']) for x in result['paths']),'images':sum(len(x['images']) for x in result['paths']),'job':job,'download':'/download/'+z.name}
                self.send_json({'summary':summary,'download':'/download/'+z.name})
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        if self.path=='/api/batch_extract':
            try:
                n=int(self.headers.get('Content-Length','0')); body=json.loads(self.rfile.read(n))
                pdfs=body.get('pdfs') or []; paths=body.get('paths') or []; text_mode=body.get('text_mode','merge')
                if not pdfs: raise ValueError('Pilih minimal satu PDF untuk batch extract')
                if not paths: raise ValueError('Paths belum didefinisikan')
                batch_job='batch_'+uuid.uuid4().hex[:10]
                combined,z=batch_extract(pdfs,paths,text_mode,batch_job)
                ok=[f for f in combined['files'] if 'error' not in f]
                failed=[f for f in combined['files'] if 'error' in f]
                summary={
                    'total_pdfs':len(pdfs),
                    'success':len(ok),
                    'failed':len(failed),
                    'errors':[{'source':f['source'],'error':f['error']} for f in failed],
                    'text_lines':sum(len(x['text']) for f in ok for x in f['paths']),
                    'images':sum(len(x['images']) for f in ok for x in f['paths']),
                    'job':batch_job,
                    'download':'/download/'+z.name,
                }
                self.send_json({'summary':summary,'download':'/download/'+z.name})
            except Exception as e: self.send_json({'error':str(e)},500)
            return
        self.send_response(404); self.end_headers()

if __name__=='__main__':
    server=ThreadingHTTPServer(('127.0.0.1',PORT),Handler)
    url=f'http://127.0.0.1:{PORT}/'
    print('PDF Path Extractor:',url)
    threading.Timer(0.8,lambda:webbrowser.open(url)).start()
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close()
